"""真实模板/CLI/XLSX/XML 回归；仅构造无害合成名称，不计算公式。"""
import importlib.util
import json
import subprocess
import sys
import tempfile
import unittest
import zipfile
import xml.etree.ElementTree as ET
from pathlib import Path

SCRIPT = Path(__file__).with_name('script.py')
TEMPLATE = SCRIPT.parent.parent / 'templates' / '导入商品信息.xlsx'
HAS_OPENPYXL = importlib.util.find_spec('openpyxl') is not None

@unittest.skipUnless(HAS_OPENPYXL, '需要已有 openpyxl；不自动安装')
class LiteralWorkbookTests(unittest.TestCase):
    def run_cli(self, items, output):
        return subprocess.run([sys.executable, str(SCRIPT), '--output', str(output)],
            input=json.dumps(items, ensure_ascii=False), text=True, capture_output=True)

    def test_actual_workbook_keeps_names_literal_and_template_contract(self):
        from openpyxl import load_workbook
        names = ['=1+1', '+1+1', '-1+1', '@SUM(A1)', ' =1+1 ', '\t=1+1',
                 ' 计算机软件（已录制） ', '中文与😀', '首行\n次行', '0' * 32767]
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / 'goods.xlsx'
            items = [{'类别':9, '类似群':'0901', '商品名称':s} for s in names]
            result = self.run_cli(items, out)
            self.assertEqual(result.returncode, 0, result.stderr)
            wb = load_workbook(out, data_only=False)
            original = load_workbook(TEMPLATE)
            self.assertEqual(wb.sheetnames, ['导入商品信息'])
            ws = wb.active
            self.assertEqual(ws.sheet_state, 'visible')
            self.assertEqual(ws.max_column, 4)
            self.assertEqual(ws.max_row, len(items) + 1)
            self.assertEqual([c.value for c in ws[1]], ['序号','商品类别','类似群','商品名称'])
            for col in 'ABCD':
                self.assertEqual(ws.column_dimensions[col].width, original.active.column_dimensions[col].width)
                self.assertEqual(ws[f'{col}1']._style, original.active[f'{col}1']._style)
            for i, name in enumerate(names, 2):
                self.assertEqual([(ws.cell(i,c).value,ws.cell(i,c).data_type) for c in range(1,5)],
                    [(i-1,'n'),(9,'n'),('0901','s'),(name,'s')])
            self.assertFalse(wb._external_links)
            self.assertFalse(list(wb.defined_names))
            ns = {'s':'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
            with zipfile.ZipFile(out) as z:
                self.assertFalse(any('vba' in n.lower() or 'externallink' in n.lower() for n in z.namelist()))
                for n in z.namelist():
                    if n.startswith('xl/worksheets/') and n.endswith('.xml'):
                        xml = ET.fromstring(z.read(n))
                        self.assertEqual(xml.findall('.//s:f', ns), [])
                xml = ET.fromstring(z.read('xl/worksheets/sheet1.xml'))
                c = xml.find(".//s:c[@r='D2']", ns)
                self.assertEqual(c.get('t'), 'inlineStr')
                self.assertEqual(c.find('s:is/s:t', ns).text, '=1+1')

    def test_overlength_name_fails_without_truncation_or_output(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / 'goods.xlsx'
            result = self.run_cli([{'类别':9,'类似群':'0901','商品名称':'长'*32768}], out)
            self.assertNotEqual(result.returncode, 0)
            self.assertIn('32767', result.stderr)
            self.assertFalse(out.exists())

    def test_whitespace_only_name_rejected_and_existing_file_untouched(self):
        with tempfile.TemporaryDirectory() as d:
            out = Path(d) / 'goods.xlsx'
            result = self.run_cli([{'类别':9,'类似群':'0901','商品名称':' \t\n'}], out)
            self.assertNotEqual(result.returncode, 0)
            self.assertFalse(out.exists())
            out.write_bytes(b'SENTINEL')
            result = self.run_cli([{'类别':9,'类似群':'0901','商品名称':'=1+1'}], out)
            self.assertNotEqual(result.returncode, 0)
            self.assertEqual(out.read_bytes(), b'SENTINEL')

if __name__ == '__main__':
    unittest.main()
